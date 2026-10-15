"""Stage 1: deterministic text extraction for every claim folder.

Writes cache/text/<tracking>.json with per-file text + a flag for files that
need vision/OCR (scanned PDFs, images). Safe to re-run; skips cached folders.
"""
import json
import os
import subprocess
import sys
from concurrent.futures import ProcessPoolExecutor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS = os.path.join(ROOT, "Security Deposit Itemization Documents")
OUT = os.path.join(ROOT, "cache", "text")
IMAGE_EXT = {".jpg", ".jpeg", ".png", ".gif", ".jfif", ".heic", ".webp"}
MIN_TEXT_CHARS = 80  # below this a PDF page set is treated as scanned


def pdf_text(path):
    try:
        r = subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, timeout=120)
        return r.stdout.decode("utf-8", "replace")
    except Exception as e:  # noqa: BLE001
        return f"<<pdftotext error: {e}>>"


def pdf_pages(path):
    try:
        r = subprocess.run(["pdfinfo", path], capture_output=True, timeout=60)
        for line in r.stdout.decode("utf-8", "replace").splitlines():
            if line.startswith("Pages:"):
                return int(line.split()[1])
    except Exception:  # noqa: BLE001
        pass
    return None


def docx_text(path):
    try:
        import docx  # python-docx
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    except Exception as e:  # noqa: BLE001
        return f"<<docx error: {e}>>"


def xlsx_text(path):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"## sheet {ws.title}")
            for row in ws.iter_rows(values_only=True):
                if any(v is not None for v in row):
                    parts.append(" | ".join("" if v is None else str(v) for v in row))
        return "\n".join(parts)
    except Exception as e:  # noqa: BLE001
        return f"<<xlsx error: {e}>>"


def process(tracking):
    out_path = os.path.join(OUT, f"{tracking}.json")
    if os.path.exists(out_path):
        return tracking, "cached"
    folder = os.path.join(DOCS, tracking)
    files = []
    for dirpath, _, names in os.walk(folder):
        for n in sorted(names):
            if n.startswith("."):
                continue
            p = os.path.join(dirpath, n)
            rel = os.path.relpath(p, folder)
            ext = os.path.splitext(n)[1].lower()
            rec = {"file": rel, "ext": ext, "bytes": os.path.getsize(p)}
            if ext == ".pdf":
                t = pdf_text(p)
                rec["pages"] = pdf_pages(p)
                rec["text"] = t
                rec["needs_vision"] = len(t.strip()) < MIN_TEXT_CHARS
            elif ext == ".docx":
                rec["text"] = docx_text(p)
                rec["needs_vision"] = False
            elif ext in (".xlsx",):
                rec["text"] = xlsx_text(p)
                rec["needs_vision"] = False
            elif ext in (".txt", ".rmd"):
                with open(p, "r", errors="replace") as f:
                    rec["text"] = f.read()
                rec["needs_vision"] = False
            elif ext in IMAGE_EXT:
                rec["text"] = ""
                rec["needs_vision"] = True
            else:
                rec["text"] = ""
                rec["needs_vision"] = ext not in (".zip", ".msg")
                rec["unsupported"] = True
            files.append(rec)
    with open(out_path, "w") as f:
        json.dump({"tracking": int(tracking), "files": files}, f)
    return tracking, len(files)


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    folders = sorted((d for d in os.listdir(DOCS) if d.isdigit()), key=int)
    if len(sys.argv) > 1:
        folders = [f for f in folders if f in set(sys.argv[1:])]
    with ProcessPoolExecutor(max_workers=8) as ex:
        for i, (t, n) in enumerate(ex.map(process, folders), 1):
            if i % 100 == 0:
                print(i, t, n, flush=True)
    print("done", len(folders))
